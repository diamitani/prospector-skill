#!/usr/bin/env python3
"""Prospector Core: turn a prospect list into a SURE list.

This script is the deterministic judge. It never trusts a claim it did not check itself.
An email reaches the SURE list only when the domain has a real MX record AND either:
  (a) published  - the exact address is printed on the company's own site, next to the
                   person's full name, and its local part matches that name; or
  (b) verified   - an approved mailbox verifier returned `valid` (not catch-all) for the
                   exact address, and the person's name was found on a page this script fetched.
Everything else goes to REVIEW (with the best unverified candidate, labelled) or REJECT.

Usage:
  python3 verify_list.py prospects.csv --out results/
  python3 verify_list.py prospects.xlsx --out results/ --verifier hunter --approved-addresses approved.txt
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import ipaddress
import json
import os
import re
import socket
import sys
import threading
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
import urllib.robotparser
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path

VERSION = '0.4.0'
UA = f'ProspectorCore/{VERSION} (contact verification; respects robots.txt)'

EMAIL_RE = re.compile(
    r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9](?:[A-Za-z0-9\-]*[A-Za-z0-9])?"
    r"(?:\.[A-Za-z0-9](?:[A-Za-z0-9\-]*[A-Za-z0-9])?)*\.[A-Za-z]{2,24}")
DOMAIN_RE = re.compile(r'(?:[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?\.)+[a-z]{2,24}')
WORD_RE = re.compile(r"[^\W\d_][^\W\d_'’\-]*")

WEBMAIL = {'gmail.com', 'googlemail.com', 'yahoo.com', 'ymail.com', 'outlook.com', 'hotmail.com',
           'live.com', 'msn.com', 'icloud.com', 'me.com', 'mac.com', 'aol.com', 'proton.me',
           'protonmail.com', 'gmx.com', 'gmx.net', 'mail.com', 'yandex.com', 'yandex.ru', 'qq.com',
           '163.com', 'comcast.net', 'verizon.net', 'att.net', 'hey.com'}
DISPOSABLE = {'mailinator.com', 'guerrillamail.com', '10minutemail.com', 'tempmail.com', 'temp-mail.org',
              'yopmail.com', 'trashmail.com', 'sharklasers.com', 'getnada.com', 'dispostable.com'}
ROLE = {'info', 'sales', 'support', 'contact', 'hello', 'hi', 'admin', 'office', 'team', 'help',
        'billing', 'marketing', 'press', 'media', 'pr', 'careers', 'jobs', 'hr', 'hiring', 'recruiting',
        'noreply', 'no-reply', 'donotreply', 'webmaster', 'postmaster', 'hostmaster', 'abuse',
        'enquiries', 'inquiries', 'enquiry', 'service', 'services', 'accounts', 'accounting', 'finance',
        'legal', 'privacy', 'security', 'mail', 'general', 'reception', 'partners', 'partnerships',
        'investors', 'ir', 'events', 'news', 'feedback', 'orders', 'customerservice', 'all'}
NAME_STOP = {'mr', 'mrs', 'ms', 'miss', 'dr', 'prof', 'jr', 'sr', 'ii', 'iii', 'iv', 'phd', 'mba',
             'cpa', 'md', 'esq'}
COMPANY_STOP = {'inc', 'llc', 'ltd', 'limited', 'corp', 'corporation', 'co', 'company', 'gmbh', 'plc',
                'sa', 'ag', 'bv', 'the', 'group', 'holdings', 'pty', 'srl', 'llp', 'lp'}
BLOCKED_HOSTS = ('linkedin.com', 'lnkd.in')  # never fetched: LinkedIn prohibits automated access
CRAWL_PATHS = ['/', '/contact', '/contact-us', '/about', '/about-us', '/team', '/our-team',
               '/leadership', '/people', '/company', '/press', '/staff']
LINK_HINT = re.compile(r'team|about|people|leadership|contact|press|staff|management|who-we-are|founders', re.I)

# Email patterns the syntax finder knows. render() turns one into a local part for a person.
PATTERNS = ['first.last', 'flast', 'first', 'firstlast', 'first_last', 'firstl', 'f.last',
            'last.first', 'lastf', 'last']
DEFAULT_GUESSES = ['first.last', 'flast', 'first']
WINDOW = 160  # characters either side of an address searched for the person's name

DNS_REJECT = {'null_mx', 'nxdomain', 'no_route'}


def now():
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def fold(value):
    """Accent-insensitive, case-insensitive, punctuation-free text for matching."""
    text = unicodedata.normalize('NFKD', str(value or ''))
    text = ''.join(c for c in text if not unicodedata.combining(c)).casefold()
    return re.sub(r'[^a-z0-9]+', ' ', text).strip()


def clean_domain(value):
    value = str(value or '').strip().lower()
    if not value:
        return None
    if '@' in value and '://' not in value:
        value = value.rsplit('@', 1)[1]
    if '://' not in value:
        value = 'http://' + value
    try:
        host = (urllib.parse.urlsplit(value).hostname or '').rstrip('.')
    except ValueError:
        return None
    if host.startswith('www.'):
        host = host[4:]
    return host if DOMAIN_RE.fullmatch(host) else None


def syntax_ok(address):
    if not isinstance(address, str) or len(address) > 254 or address.count('@') != 1:
        return False
    local, domain = address.rsplit('@', 1)
    return bool(local and len(local) <= 64 and not local.startswith('.') and not local.endswith('.')
                and '..' not in local and re.fullmatch(r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+", local)
                and DOMAIN_RE.fullmatch(domain.lower()))


def related(a, b):
    """Same organisation domain: equal, or one is a subdomain of the other."""
    return bool(a and b) and (a == b or a.endswith('.' + b) or b.endswith('.' + a))


def base_local(local):
    return local.lower().split('+', 1)[0]


ROLE_COMPACT = {r.replace('-', '') for r in ROLE}


def is_role(local):
    return re.sub(r'[._\-]', '', base_local(local)) in ROLE_COMPACT


def host_of(url):
    try:
        return (urllib.parse.urlsplit(url).hostname or '').lower().rstrip('.')
    except ValueError:
        return ''


def is_blocked_host(host):
    return any(host == b or host.endswith('.' + b) for b in BLOCKED_HOSTS)


# ---------------------------------------------------------------- names and patterns

def parse_name(name=None, first=None, last=None):
    """Return {first, last, last_full, display} or None when the name cannot be parsed safely."""
    if first and last:
        f_tokens, l_tokens = fold(first).split(), [t for t in fold(last).split() if t not in NAME_STOP]
        display = f'{first} {last}'.strip()
    else:
        raw = str(name or '').strip()
        if ',' in raw:
            head, tail = raw.split(',', 1)
            if fold(tail) and fold(tail).split()[0] not in NAME_STOP:
                raw = f'{tail} {head}'
        tokens = [t for t in fold(raw).split() if t not in NAME_STOP]
        if len(tokens) < 2:
            return None
        f_tokens, l_tokens, display = tokens[:1], tokens[1:], str(name).strip()
        l_tokens = l_tokens[-1:]  # without separate columns, particles are ambiguous
    if not f_tokens or not l_tokens or len(f_tokens[0]) < 2 or len(l_tokens[-1]) < 2:
        return None
    return {'first': f_tokens[0], 'last': l_tokens[-1], 'last_full': ''.join(l_tokens), 'display': display}


def person_from_words(a, b):
    return parse_name(first=a, last=b)


def render(pattern, person):
    f, l = person['first'], person['last_full']
    return {'first.last': f'{f}.{l}', 'flast': f[0] + l, 'first': f, 'firstlast': f + l,
            'first_last': f'{f}_{l}', 'firstl': f + l[0], 'f.last': f'{f[0]}.{l}',
            'last.first': f'{l}.{f}', 'lastf': l + f[0], 'last': l}[pattern]


def local_matches(local, person):
    """True when an address local part is a plausible rendering of the person's name."""
    compact = re.sub(r'[._\-]', '', base_local(local))
    f = person['first']
    variants = {f}
    for l in {person['last'], person['last_full']}:
        variants |= {l, f + l, l + f, f[0] + l, f + l[0], l + f[0]}
    return compact in variants


# ---------------------------------------------------------------- html

def html_to_text(raw):
    t = re.sub(r'(?is)<(script|style|noscript|template|svg)\b.*?</\1\s*>', ' ', raw)
    t = re.sub(r'(?s)<!--.*?-->', ' ', t)
    t = re.sub(r'(?is)<a\b[^>]*?href\s*=\s*["\']mailto:([^"\'?>\s]+)[^"\']*["\'][^>]*>(.*?)</a\s*>',
               lambda m: f' {m.group(2)} {urllib.parse.unquote(m.group(1))} ', t)
    t = re.sub(r'(?i)<(br|hr|/p|/div|/li|/tr|/td|/h[1-6]|/section|/article|/header|/footer)\b[^>]*>', '\n', t)
    t = re.sub(r'(?s)<[^>]+>', ' ', t)
    t = html.unescape(t)
    t = re.sub(r'\s*[\[\(\{]\s*at\s*[\]\)\}]\s*', '@', t, flags=re.I)
    t = re.sub(r'\s*[\[\(\{]\s*dot\s*[\]\)\}]\s*', '.', t, flags=re.I)
    t = re.sub(r'[ \t\r\f\v]+', ' ', t)
    return re.sub(r'\n\s*\n+', '\n', t)


def discover_links(raw, page_url, domain):
    out = []
    for href in re.findall(r'(?is)<a\b[^>]*?href\s*=\s*["\']([^"\'#]+)["\']', raw):
        url = urllib.parse.urljoin(page_url, href.strip())
        if url.startswith(('http://', 'https://')) and related(host_of(url).removeprefix('www.'), domain) \
                and LINK_HINT.search(urllib.parse.urlsplit(url).path or ''):
            out.append(url.split('#')[0])
    return list(dict.fromkeys(out))


# ---------------------------------------------------------------- fetching

class Blocked(Exception):
    pass


class _NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        return None


class SafeFetcher:
    """Public-web GET only: http(s), no private/loopback IPs, robots.txt respected, no LinkedIn."""

    def __init__(self, timeout=10, max_bytes=1_500_000, delay=0.4, allow_private=False):
        self.timeout, self.max_bytes, self.delay, self.allow_private = timeout, max_bytes, delay, allow_private
        self.opener = urllib.request.build_opener(_NoRedirect)
        self.robots, self.last_hit, self.lock = {}, {}, threading.Lock()

    def _check_host(self, host):
        if not host or is_blocked_host(host):
            raise Blocked('blocked_host')
        if self.allow_private:
            return
        for info in socket.getaddrinfo(host, None):
            ip = ipaddress.ip_address(info[4][0].split('%')[0])
            if not ip.is_global:
                raise Blocked('non_public_address')

    def _throttle(self, host):
        with self.lock:
            wait = self.last_hit.get(host, 0) + self.delay - time.monotonic()
            self.last_hit[host] = max(time.monotonic(), self.last_hit.get(host, 0) + self.delay)
        if wait > 0:
            time.sleep(wait)

    def _get(self, url):
        """One request, no redirects. Returns (status, headers, body_bytes, location)."""
        parts = urllib.parse.urlsplit(url)
        if parts.scheme not in ('http', 'https'):
            raise Blocked('bad_scheme')
        self._check_host(parts.hostname)
        self._throttle(parts.hostname)
        req = urllib.request.Request(url, headers={'User-Agent': UA, 'Accept': 'text/html,text/plain;q=0.9'})
        try:
            with self.opener.open(req, timeout=self.timeout) as r:
                return r.status, r.headers, r.read(self.max_bytes), None
        except urllib.error.HTTPError as e:
            return e.code, e.headers, b'', e.headers.get('Location') if e.headers else None

    def _robots_ok(self, url):
        parts = urllib.parse.urlsplit(url)
        key = f'{parts.scheme}://{parts.netloc}'
        with self.lock:
            rp = self.robots.get(key)
        if rp is None:
            rp = urllib.robotparser.RobotFileParser()
            try:
                robots_url = key + '/robots.txt'
                for _ in range(4):  # robots.txt often redirects (e.g. apex -> www)
                    status, _, body, location = self._get(robots_url)
                    if status in (301, 302, 303, 307, 308) and location:
                        robots_url = urllib.parse.urljoin(robots_url, location)
                        continue
                    break
                if status == 200:
                    rp.parse(body.decode('utf-8', 'replace').splitlines())
                elif 400 <= status < 500:
                    rp.allow_all = True
                else:
                    rp.disallow_all = True
            except Exception:
                rp.disallow_all = True
            with self.lock:
                self.robots[key] = rp
        return rp.can_fetch(UA, url)

    def fetch(self, url):
        result = {'ok': False, 'url': url, 'final_url': url, 'status': None, 'text': '', 'raw': '',
                  'error': None, 'retrieved_at': now(), 'sha256': None}
        try:
            for _ in range(4):
                parts = urllib.parse.urlsplit(url)
                if parts.scheme not in ('http', 'https'):
                    raise Blocked('bad_scheme')
                self._check_host(parts.hostname)
                if not self._robots_ok(url):
                    raise Blocked('robots_disallow')
                status, headers, body, location = self._get(url)
                result['status'] = status
                if status in (301, 302, 303, 307, 308) and location:
                    url = urllib.parse.urljoin(url, location)
                    result['final_url'] = url
                    continue
                if status != 200:
                    result['error'] = f'http_{status}'
                    return result
                ctype = (headers.get('Content-Type') or '').lower()
                if ctype and not any(t in ctype for t in ('text/html', 'text/plain', 'xhtml')):
                    result['error'] = 'not_text'
                    return result
                charset = headers.get_content_charset() or 'utf-8'
                raw = body.decode(charset, 'replace')
                result.update(ok=True, raw=raw, text=html_to_text(raw) if 'plain' not in ctype else raw,
                              sha256=hashlib.sha256(body).hexdigest())
                return result
            result['error'] = 'too_many_redirects'
        except Blocked as e:
            result['error'] = str(e)
        except Exception as e:  # network errors are data, not crashes
            result['error'] = type(e).__name__
        return result


# ---------------------------------------------------------------- DNS

class DnsPythonResolver:
    name = 'dnspython'

    def __init__(self, timeout=6):
        import dns.resolver  # noqa: F401  (ImportError selects the DoH fallback)
        self.resolver = dns.resolver.Resolver()
        self.resolver.lifetime = timeout

    def status(self, domain):
        import dns.exception
        import dns.resolver
        try:
            answer = self.resolver.resolve(domain, 'MX')
            if any(r.exchange.to_text() == '.' for r in answer):
                return 'null_mx'
            return 'mx_present'
        except dns.resolver.NXDOMAIN:
            return 'nxdomain'
        except dns.resolver.NoAnswer:
            for rtype in ('A', 'AAAA'):
                try:
                    self.resolver.resolve(domain, rtype)
                    return 'implicit_mx_possible'
                except (dns.resolver.NoAnswer, dns.resolver.NXDOMAIN):
                    continue
                except dns.exception.DNSException:
                    return 'unknown'
            return 'no_route'
        except dns.exception.DNSException:
            return 'unknown'


class DohResolver:
    """DNS-over-HTTPS fallback (Google public DNS JSON API) when dnspython is not installed."""
    name = 'doh:dns.google'

    def __init__(self, timeout=8):
        self.timeout = timeout

    def _query(self, name, rtype):
        url = 'https://dns.google/resolve?' + urllib.parse.urlencode({'name': name, 'type': rtype})
        req = urllib.request.Request(url, headers={'Accept': 'application/dns-json', 'User-Agent': UA})
        with urllib.request.urlopen(req, timeout=self.timeout) as r:
            return json.load(r)

    def status(self, domain):
        try:
            j = self._query(domain, 'MX')
            if j.get('Status') == 3:
                return 'nxdomain'
            if j.get('Status') != 0:
                return 'unknown'
            mx = [a.get('data', '') for a in j.get('Answer', []) if a.get('type') == 15]
            if mx:
                return 'null_mx' if any(re.fullmatch(r'\d+\s+\.', d.strip()) for d in mx) else 'mx_present'
            for rtype, code in (('A', 1), ('AAAA', 28)):
                j = self._query(domain, rtype)
                if j.get('Status') == 0 and any(a.get('type') == code for a in j.get('Answer', [])):
                    return 'implicit_mx_possible'
            return 'no_route'
        except Exception:
            return 'unknown'


def make_resolver():
    try:
        return DnsPythonResolver()
    except ImportError:
        return DohResolver()


# ---------------------------------------------------------------- verifier

VERIFIER_STATUSES = {'valid', 'invalid', 'accept_all', 'webmail', 'disposable', 'unknown'}


class HunterVerifier:
    """Hunter Email Verifier (https://hunter.io/api-documentation). Key from HUNTER_API_KEY only.
    Not live-tested in this package: confirm the endpoint, auth header and your free quota first."""
    name = 'hunter'

    def __init__(self, timeout=30):
        self.key = os.environ.get('HUNTER_API_KEY')
        if not self.key:
            raise SystemExit('HUNTER_API_KEY is not set; refusing to run the verifier.')
        self.timeout = timeout

    def verify(self, email):
        url = 'https://api.hunter.io/v2/email-verifier?' + urllib.parse.urlencode({'email': email})
        req = urllib.request.Request(url, headers={'X-API-KEY': self.key, 'User-Agent': UA})
        try:
            with urllib.request.urlopen(req, timeout=self.timeout) as r:
                data = json.load(r).get('data', {})
        except urllib.error.HTTPError as e:
            return {'status': 'quota_exhausted' if e.code in (402, 429) else 'provider_error', 'raw_status': f'http_{e.code}'}
        raw = str(data.get('status', 'unknown'))
        status = raw if raw in VERIFIER_STATUSES else 'unknown'
        if data.get('accept_all') and status == 'valid':
            status = 'accept_all'
        return {'status': status, 'raw_status': raw, 'score': data.get('score')}


class VerifierGate:
    """Sends an address to the third-party verifier only if it is on the approved list,
    within the call budget. Every requested address is recorded for the next approval round."""

    def __init__(self, inner=None, approved=None, max_calls=20):
        self.inner, self.approved, self.max_calls = inner, set(approved or ()), max_calls
        self.calls, self.requested, self.cache, self.lock = 0, [], {}, threading.Lock()

    def verify(self, email):
        base = {'provider': getattr(self.inner, 'name', None), 'checked_at': None, 'raw_status': None}
        if self.inner is None:
            return dict(base, status='not_configured')
        with self.lock:
            if email in self.cache:
                return self.cache[email]
            if email not in self.requested:
                self.requested.append(email)
            if email not in self.approved:
                return dict(base, status='approval_required')
            if self.calls >= self.max_calls:
                return dict(base, status='quota_exhausted')
            self.calls += 1
        try:
            out = self.inner.verify(email)
        except Exception as e:
            out = {'status': 'provider_error', 'raw_status': type(e).__name__}
        status = out.get('status')
        if status not in VERIFIER_STATUSES | {'quota_exhausted', 'provider_error'}:
            status = 'unknown'
        result = dict(base, status=status, raw_status=out.get('raw_status'), checked_at=now(),
                      score=out.get('score'))
        with self.lock:
            self.cache[email] = result
        return result


# ---------------------------------------------------------------- input

ALIASES = {
    'name': ['name', 'full name', 'fullname', 'contact', 'contact name', 'person', 'person name'],
    'first_name': ['first name', 'firstname', 'first', 'given name'],
    'last_name': ['last name', 'lastname', 'last', 'surname', 'family name'],
    'title': ['title', 'job title', 'role', 'position'],
    'company': ['company', 'company name', 'organization', 'organisation', 'account', 'account name', 'employer'],
    'domain': ['domain', 'company domain', 'website', 'company website', 'url', 'site', 'web', 'homepage'],
    'email': ['email', 'email address', 'e mail', 'work email', 'business email'],
    'person_source_urls': ['person source urls', 'person source url', 'person sources', 'source url',
                           'source urls', 'sources', 'source', 'evidence', 'evidence url', 'evidence urls'],
    'linkedin_url': ['linkedin', 'linkedin url', 'linkedin profile', 'profile url', 'person linkedin url'],
    'signal': ['signal', 'buying signal', 'trigger'],
    'signal_url': ['signal url', 'signal source', 'signal source url'],
}
HEADER_MAP = {alias: key for key, names in ALIASES.items() for alias in names}


def read_table(path):
    path = Path(path)
    if path.suffix.lower() in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = [['' if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
        header, body = rows[0], rows[1:]
        return [dict(zip(header, r)) for r in body if any(c.strip() for c in r)]
    text = path.read_text(encoding='utf-8-sig')
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel
    return [r for r in csv.DictReader(text.splitlines(), dialect=dialect) if any((v or '').strip() for v in r.values())]


def split_urls(value):
    return [u for u in re.split(r'[\s;|,]+', str(value or '')) if u.startswith(('http://', 'https://'))]


def normalise_row(raw):
    row = {k: '' for k in ALIASES}
    extra = {}
    for header, value in raw.items():
        key = HEADER_MAP.get(fold(header or ''))
        value = '' if value is None else str(value).strip()
        if key and not row[key]:
            row[key] = value
        elif header:
            extra[header] = value
    urls = split_urls(row['person_source_urls'])
    li = [u for u in urls if is_blocked_host(host_of(u))]
    row['person_source_urls'] = [u for u in urls if u not in li]
    row['linkedin_url'] = row['linkedin_url'] or (li[0] if li else '')
    row['extra'] = extra
    return row


def load_list(path):
    entries = []
    for line in Path(path).read_text(encoding='utf-8-sig').splitlines():
        if line.lstrip().startswith('#'):
            continue
        entries += [t.strip().lower() for t in re.split(r'[\s,;]+', line) if t.strip()]
    return entries


# ---------------------------------------------------------------- the judge

class Judge:
    def __init__(self, fetcher, resolver, gate, crawl=True, max_pages=8, evidence_policy='company_site',
                 suppression=(), workers=6):
        self.fetcher, self.resolver, self.gate = fetcher, resolver, gate
        self.crawl_on, self.max_pages, self.policy, self.workers = crawl, max_pages, evidence_policy, workers
        sup = {s.lower() for s in suppression}
        self.sup_emails = {s for s in sup if '@' in s}
        self.sup_domains = {clean_domain(s) for s in sup if '@' not in s} - {None}
        self.url_cache, self.dns_cache, self.crawl_cache, self.pattern_cache = {}, {}, {}, {}
        self.lock = threading.Lock()

    # caches -------------------------------------------------------------
    def page(self, url, domain=None):
        with self.lock:
            hit = self.url_cache.get(url)
        if hit is None:
            hit = self.fetcher.fetch(url)
            with self.lock:
                self.url_cache[url] = hit
        return hit

    def tagged(self, fetched, domain):
        final_host = host_of(fetched['final_url']).removeprefix('www.')
        return {'url': fetched['final_url'], 'text': fetched['text'], 'retrieved_at': fetched['retrieved_at'],
                'sha256': fetched['sha256'],
                'host_class': 'company_site' if related(final_host, domain) else 'third_party'}

    def dns(self, domain):
        with self.lock:
            if domain in self.dns_cache:
                return self.dns_cache[domain]
        result = {'status': self.resolver.status(domain), 'checked_at': now(),
                  'resolver': getattr(self.resolver, 'name', 'custom')}
        with self.lock:
            self.dns_cache[domain] = result
        return result

    def crawl(self, domain):
        with self.lock:
            if domain in self.crawl_cache:
                return self.crawl_cache[domain]
        pages, attempts, queue, seen = [], 0, [], set()
        home = self.page(f'https://{domain}/')
        if not home['ok']:
            home = self.page(f'https://www.{domain}/')
        base = f'https://{domain}'
        if home['ok']:
            attempts += 1
            final = urllib.parse.urlsplit(home['final_url'])
            if related(host_of(home['final_url']).removeprefix('www.'), domain):
                base = f'{final.scheme}://{final.netloc}'
            pages.append(self.tagged(home, domain))
            seen.add(home['final_url'].rstrip('/'))
            queue += discover_links(home.get('raw', ''), home['final_url'], domain)
        else:
            attempts += 1  # the failed homepage still costs budget, but keep trying the usual pages
        queue += [base + p for p in CRAWL_PATHS[1:]]
        for url in dict.fromkeys(queue):
            if attempts >= self.max_pages:
                break
            if url.rstrip('/') in seen:
                continue
            attempts += 1
            got = self.page(url)
            seen.add(url.rstrip('/'))
            if got['ok'] and got['final_url'].rstrip('/') not in {p['url'].rstrip('/') for p in pages}:
                pages.append(self.tagged(got, domain))
                seen.add(got['final_url'].rstrip('/'))
        with self.lock:
            self.crawl_cache[domain] = pages
        return pages

    def patterns(self, domain):
        """Syntax finder: learn the company's address format from named employees on its own site."""
        with self.lock:
            if domain in self.pattern_cache:
                return self.pattern_cache[domain]
        votes = {}
        for page in self.crawl_cache.get(domain, []):
            if page['host_class'] != 'company_site':
                continue
            text = page['text']
            for m in EMAIL_RE.finditer(text):
                address = m.group(0).lower()
                local, dom = address.rsplit('@', 1)
                if not related(dom, domain) or is_role(local):
                    continue
                window = EMAIL_RE.sub(' ', text[max(0, m.start() - 120):m.start()] + ' ' + text[m.end():m.end() + 120])
                words = WORD_RE.findall(window)
                matched = set()
                for a, b in zip(words, words[1:]):
                    if a[:1].isupper() and b[:1].isupper():
                        p = person_from_words(a, b)
                        if p:
                            matched |= {pat for pat in PATTERNS if render(pat, p) == base_local(local)}
                if len(matched) == 1:
                    votes.setdefault(matched.pop(), set()).add(address)
        ranked = sorted(votes.items(), key=lambda kv: -len(kv[1]))
        result = {'status': 'none', 'pattern': None, 'examples': []}
        if ranked:
            top, examples = ranked[0]
            runner_up = len(ranked[1][1]) if len(ranked) > 1 else 0
            if len(examples) == runner_up:
                status = 'ambiguous'
            elif len(examples) >= 2:
                status = 'corroborated'
            else:
                status = 'weak'
            result = {'status': status, 'pattern': top, 'examples': sorted(examples)[:5],
                      'all_votes': {k: len(v) for k, v in ranked}}
        with self.lock:
            self.pattern_cache[domain] = result
        return result

    def prewarm(self, rows):
        domains = {r['_domain'] for r in rows if r.get('_domain')}
        urls = {u for r in rows for u in r['person_source_urls'] + split_urls(r['signal_url'])}
        with ThreadPoolExecutor(self.workers) as pool:
            list(pool.map(self.dns, domains))
            if self.crawl_on:
                list(pool.map(self.crawl, domains))
            list(pool.map(self.page, urls))
        for d in domains:
            self.patterns(d)

    # checks ---------------------------------------------------------------
    def suppressed(self, domain=None, email=None):
        return (email or '') in self.sup_emails or any(related(domain or '', d) for d in self.sup_domains) or \
            bool(email and any(related(email.rsplit('@', 1)[1], d) for d in self.sup_domains))

    def confirm_person(self, person, company, domain, pages):
        pattern = rf"\b{re.escape(person['first'])}\b(?: [a-z0-9]+){{0,2}} {re.escape(person['last'])}\b"
        core = ' '.join(t for t in fold(company).split() if t not in COMPANY_STOP)
        for page in pages:
            text = fold(page['text'])
            if not re.search(pattern, text):
                continue
            if page['host_class'] == 'company_site' or (core and re.search(rf'\b{re.escape(core)}\b', text)) \
                    or domain in page['text'].lower():
                return page['url']
        return None

    def find_published(self, person, domain, pages):
        f, l = re.escape(person['first']), re.escape(person['last'])
        full_name = re.compile(rf'\b{f}\b(?: [a-z0-9]+){{0,2}} {l}\b|\b{l} {f}\b')
        for_me = (person['first'], person['last'])
        hits = []
        for page in pages:
            if self.policy == 'company_site' and page['host_class'] != 'company_site':
                continue
            text = page['text']
            for m in EMAIL_RE.finditer(text):
                address = m.group(0).lower()
                local, dom = address.rsplit('@', 1)
                if not related(dom, domain) or is_role(local) or not local_matches(local, person):
                    continue
                window_raw = EMAIL_RE.sub(' ', text[max(0, m.start() - WINDOW):m.start()] + ' '
                                          + text[m.end():m.end() + WINDOW])
                if not full_name.search(fold(window_raw)):
                    continue
                # Ambiguity guard: another nearby person whose name also fits this address.
                words = WORD_RE.findall(window_raw)
                rivals = [p for a, b in zip(words, words[1:]) if a[:1].isupper() and b[:1].isupper()
                          for p in [person_from_words(a, b)]
                          if p and (p['first'], p['last']) != for_me and local_matches(local, p)]
                if not rivals:
                    excerpt = re.sub(r'\s+', ' ', text[max(0, m.start() - 120):m.end() + 120]).strip()
                    hits.append({'address': address, 'url': page['url'], 'excerpt': excerpt[:320],
                                 'retrieved_at': page['retrieved_at'], 'sha256': page['sha256'],
                                 'host_class': page['host_class']})
        return hits

    def gate_address(self, address, domain, R):
        """Address-level checks shared by every proof path. Returns (status, reason) or None if it passes."""
        local, dom = address.rsplit('@', 1)
        if is_role(local):
            return 'REJECT', 'role_mailbox'
        if dom in WEBMAIL:
            return 'REJECT', 'webmail_domain'
        if dom in DISPOSABLE:
            return 'REJECT', 'disposable_domain'
        if self.suppressed(email=address):
            return 'REJECT', 'suppressed'
        if not related(dom, domain):
            return 'REVIEW', 'email_domain_differs_from_company_domain'
        R['dns'] = self.dns(dom)
        status = R['dns']['status']
        if status in DNS_REJECT:
            return 'REJECT', f'dns_{status}'
        if status == 'unknown':
            return 'REVIEW', 'dns_unknown'
        if status == 'implicit_mx_possible':
            return 'REVIEW', 'no_mx_record_implicit_only'
        return None

    def judge(self, idx, row):
        R = {'row_id': idx, 'name': row['name'] or f"{row['first_name']} {row['last_name']}".strip(),
             'title': row['title'], 'company': row['company'], 'domain': row['_domain'],
             'email_input': row['email'].lower() or None, 'linkedin_url': row['linkedin_url'] or None,
             'signal': row['signal'] or None, 'status': None, 'reasons': [], 'notes': [], 'email': None,
             'proof': None, 'ownership': None, 'evidence': None, 'dns': None, 'verifier': None,
             'best_candidate': None, 'candidate_basis': None, 'person_confirmed_url': None,
             'signal_check': None, 'pattern': None, 'checked_at': now()}

        def done(status, reason=None):
            R['status'] = status
            if reason:
                R['reasons'].append(reason)
            if status != 'SURE':  # only proven rows may carry an `email`; others use best_candidate
                R['email'] = None
            return R

        supplied, domain = R['email_input'], R['domain']
        if row['signal_url']:
            R['signal_check'] = self.check_signal(row)
        if row['_domain_note']:
            R['notes'].append(row['_domain_note'])
        if supplied and not syntax_ok(supplied):
            return done('REJECT', 'syntax_invalid')
        if not domain:
            return done('REVIEW', 'no_company_domain')
        if domain in WEBMAIL:
            return done('REJECT', 'webmail_domain')
        if self.suppressed(domain=domain, email=supplied):
            return done('REJECT', 'suppressed')
        if row['linkedin_url']:
            R['notes'].append('linkedin_url_not_fetched')
        person = parse_name(row['name'], row['first_name'], row['last_name'])
        pages = list(self.crawl_cache.get(domain, [])) + [
            self.tagged(p, domain) for p in (self.page(u) for u in row['person_source_urls']) if p['ok']]
        R['pattern'] = self.pattern_cache.get(domain)
        if person:
            R['person_confirmed_url'] = self.confirm_person(person, row['company'], domain, pages)
        else:
            R['notes'].append('name_unparsed')

        hits = self.find_published(person, domain, pages) if person else []
        distinct = sorted({h['address'] for h in hits})
        if len(distinct) > 1 and supplied not in distinct:
            R['best_candidate'], R['candidate_basis'] = ' | '.join(distinct), 'published_multiple'
            return done('REVIEW', 'multiple_published_addresses')
        if distinct and supplied and supplied not in distinct:
            R['best_candidate'], R['candidate_basis'] = distinct[0], 'published_on_site'
            return done('REVIEW', 'supplied_email_differs_from_published')

        if distinct:  # proof path (a): published next to the name
            address = supplied or distinct[0]
            R['email'], R['evidence'] = address, next(h for h in hits if h['address'] == address)
            failed = self.gate_address(address, domain, R)
            if failed:
                R['best_candidate'], R['candidate_basis'] = address, 'published_on_site'
                return done(*failed)
            R['proof'] = 'published_on_company_site' if R['evidence']['host_class'] == 'company_site' \
                else 'published_on_third_party_page'
            R['ownership'] = 'source_asserted'
            return done('SURE')

        # proof path (b): mailbox verifier, for a supplied address or labelled guesses
        if supplied:
            candidates, basis = [supplied], 'list_supplied'
        elif person:
            candidates = self.guesses(person, R['pattern'], domain)
            basis = None
        else:
            return done('REVIEW', 'no_email_and_name_unparsed')
        if not candidates:
            return done('REVIEW', 'no_valid_guess_for_name')

        first = candidates[0]
        failed = self.gate_address(first, domain, R)
        if failed:
            R['best_candidate'], R['candidate_basis'] = first, basis or self.basis(first, person, R['pattern'])
            return done(*failed)
        R['best_candidate'], R['candidate_basis'] = first, basis or self.basis(first, person, R['pattern'])
        tried = []
        for address in candidates:
            v = self.gate.verify(address)
            tried.append({'address': address, **v})
            R['verifier'] = v
            if v['status'] == 'valid':
                R['email'] = address
                R['best_candidate'], R['candidate_basis'] = address, basis or self.basis(address, person, R['pattern'])
                if not R['person_confirmed_url']:
                    R['notes'].append('verifier_tried:' + ','.join(t['address'] for t in tried))
                    R['email'] = None
                    return done('REVIEW', 'mailbox_valid_but_person_unconfirmed')
                R['proof'] = 'mailbox_verified'
                R['ownership'] = 'list_supplied' if basis == 'list_supplied' else 'pattern_inferred'
                return done('SURE')
            if v['status'] == 'invalid':
                continue
            break
        R['notes'].append('verifier_tried:' + ','.join(f"{t['address']}={t['status']}" for t in tried))
        last = R['verifier']['status']
        if last == 'invalid':
            return done('REJECT', 'verifier_invalid') if basis == 'list_supplied' else done('REVIEW', 'all_guesses_invalid')
        if last in ('webmail', 'disposable'):
            return done('REJECT', f'verifier_{last}')
        reason = {'not_configured': 'not_verified_no_verifier', 'approval_required': 'awaiting_verification_approval',
                  'accept_all': 'catch_all_domain'}.get(last, f'verifier_{last}')
        return done('REVIEW', reason)

    @staticmethod
    def guesses(person, pattern, domain):
        """At most three labelled guesses: the corroborated site pattern first, then common defaults."""
        order = ([pattern['pattern']] if pattern and pattern.get('status') == 'corroborated' else []) + DEFAULT_GUESSES
        out = []
        for pat in order:
            address = f'{render(pat, person)}@{domain}'
            if address not in out and syntax_ok(address):
                out.append(address)
        return out[:3]

    def basis(self, address, person, pattern):
        if not person:
            return 'list_supplied'
        local = base_local(address.rsplit('@', 1)[0])
        if pattern and pattern.get('status') == 'corroborated' and render(pattern['pattern'], person) == local:
            return f"pattern:{pattern['pattern']} ({len(pattern['examples'])} named examples on site)"
        return 'default_guess (no published pattern found)'

    def check_signal(self, row):
        urls = split_urls(row['signal_url'])
        if not urls:
            return None
        if is_blocked_host(host_of(urls[0])):
            return 'linkedin_not_fetched'
        page = self.page(urls[0])
        if not page['ok']:
            return f"fetch_failed:{page['error']}"
        core = ' '.join(t for t in fold(row['company']).split() if t not in COMPANY_STOP)
        text = fold(page['text'])
        found = (core and re.search(rf'\b{re.escape(core)}\b', text)) or \
            (row['_domain'] and row['_domain'] in page['text'].lower())
        return 'company_found_on_signal_page' if found else 'company_not_found_on_signal_page'

    def run(self, raw_rows):
        rows = [normalise_row(r) for r in raw_rows]
        for r in rows:
            r['_domain'] = clean_domain(r['domain'])
            r['_domain_note'] = None
            if not r['_domain'] and r['email'] and '@' in r['email']:
                r['_domain'] = clean_domain(r['email'])
                r['_domain_note'] = 'company_domain_taken_from_email'
        self.prewarm(rows)
        results = [self.judge(i, r) for i, r in enumerate(rows, 1)]
        seen = {}
        for res in results:  # one SURE row per address
            if res['status'] == 'SURE':
                if res['email'] in seen:
                    res['status'], res['reasons'] = 'REJECT', [f"duplicate_of_row_{seen[res['email']]}"]
                    res['best_candidate'], res['candidate_basis'], res['email'] = res['email'], res['proof'], None
                else:
                    seen[res['email']] = res['row_id']
        return results


# ---------------------------------------------------------------- output

SURE_COLS = ['row_id', 'name', 'title', 'company', 'domain', 'email', 'proof', 'ownership', 'evidence_url',
             'evidence_excerpt', 'evidence_retrieved_at', 'person_confirmed_url', 'dns_status',
             'verifier_status', 'signal', 'signal_check', 'checked_at']
OTHER_COLS = ['row_id', 'name', 'title', 'company', 'domain', 'email_input', 'status', 'reasons',
              'best_candidate_UNVERIFIED', 'candidate_basis', 'person_confirmed_url', 'dns_status',
              'verifier_status', 'linkedin_url', 'signal', 'signal_check', 'notes', 'checked_at']


def flat(res):
    ev = res['evidence'] or {}
    return {**res, 'evidence_url': ev.get('url'), 'evidence_excerpt': ev.get('excerpt'),
            'evidence_retrieved_at': ev.get('retrieved_at'),
            'dns_status': (res['dns'] or {}).get('status'), 'verifier_status': (res['verifier'] or {}).get('status'),
            'reasons': '; '.join(res['reasons']), 'notes': '; '.join(res['notes']),
            'best_candidate_UNVERIFIED': res['best_candidate']}


def safe_cell(value):
    value = '' if value is None else str(value)
    return "'" + value if value[:1] in ('=', '+', '-', '@', '\t', '\r') else value


def write_csv(path, cols, rows):
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction='ignore')
        w.writeheader()
        for r in rows:
            w.writerow({c: safe_cell(r.get(c)) for c in cols})


def write_outputs(results, out, meta, gate):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    flats = [flat(r) for r in results]
    groups = {s: [f for f in flats if f['status'] == s] for s in ('SURE', 'REVIEW', 'REJECT')}
    write_csv(out / 'sure_list.csv', SURE_COLS, groups['SURE'])
    write_csv(out / 'review_list.csv', OTHER_COLS, groups['REVIEW'])
    write_csv(out / 'rejected.csv', OTHER_COLS, groups['REJECT'])
    try:
        from openpyxl import Workbook
        wb = Workbook()
        for i, (name, cols, rows) in enumerate([('Sure list', SURE_COLS, groups['SURE']),
                                                ('Review', OTHER_COLS, groups['REVIEW']),
                                                ('Rejected', OTHER_COLS, groups['REJECT'])]):
            ws = wb.active if i == 0 else wb.create_sheet()
            ws.title = name
            ws.append(cols)
            for r in rows:
                ws.append([safe_cell(r.get(c)) for c in cols])
        wb.save(out / 'prospector_results.xlsx')
    except ImportError:
        pass
    (out / 'all_results.json').write_text(json.dumps({'meta': meta, 'results': results}, indent=2))
    if gate.requested:
        pending = [a for a in gate.requested if a not in gate.approved]
        (out / 'verification_requests.txt').write_text(
            '# Addresses the verifier would send to the third-party provider. Approve by passing this\n'
            '# file (edited as you like) via --approved-addresses. Nothing here has been sent.\n'
            + '\n'.join(pending) + ('\n' if pending else ''))
    reasons = {}
    for r in results:
        if r['status'] != 'SURE':
            for reason in r['reasons']:
                reasons[reason] = reasons.get(reason, 0) + 1
    proofs = {}
    for r in groups['SURE']:
        proofs[r['proof']] = proofs.get(r['proof'], 0) + 1
    lines = [f"# Prospector Core run — {meta['finished_at']}", '',
             f"Input rows: **{len(results)}** · SURE: **{len(groups['SURE'])}** · REVIEW: {len(groups['REVIEW'])}"
             f" · REJECT: {len(groups['REJECT'])}", '',
             '## How the SURE rows are proven', '']
    lines += [f'- {k}: {v}' for k, v in sorted(proofs.items())] or ['- none']
    lines += ['', '## Why the other rows are not SURE', '']
    lines += [f'- {k}: {v}' for k, v in sorted(reasons.items(), key=lambda kv: -kv[1])] or ['- n/a']
    lines += ['', '## Run settings', '',
              f"- DNS resolver: {meta['resolver']} (preflight: {meta['dns_preflight']})",
              f"- Company-site crawl: {meta['crawl']} (max {meta['max_pages']} pages/domain), evidence policy: {meta['evidence_policy']}",
              f"- Verifier: {meta['verifier']} · calls made: {gate.calls} · awaiting approval: "
              f"{len([a for a in gate.requested if a not in gate.approved])}",
              f"- Suppression list: {meta['suppression']}", '',
              '## What SURE means (and does not)', '',
              '- SURE = the exact address is published on the company site next to the person\'s name, or an approved '
              'verifier said the mailbox exists and the person was found on a fetched page; and the domain has MX.',
              '- MX alone never makes a row SURE: it proves the domain takes mail, not that an inbox exists.',
              '- `best_candidate_UNVERIFIED` in the review list is a lead to check, not an answer.',
              '- SURE is not consent to contact. No email was sent; nothing was written to a CRM.']
    if meta['blockers']:
        lines += ['', '## Blockers', ''] + [f'- {b}' for b in meta['blockers']]
    (out / 'run_report.md').write_text('\n'.join(lines) + '\n')
    return groups


def main(argv=None):
    p = argparse.ArgumentParser(description='Turn a prospect list into a proof-backed SURE list.')
    p.add_argument('input', help='CSV or XLSX prospect list')
    p.add_argument('--out', default='prospector-output')
    p.add_argument('--no-crawl', action='store_true', help='do not read the companies\' own websites')
    p.add_argument('--max-pages', type=int, default=8, help='pages fetched per company site (default 8)')
    p.add_argument('--evidence-policy', choices=['company_site', 'any'], default='company_site',
                   help='where a published address may be found (default: company site only)')
    p.add_argument('--suppress', help='file of emails/domains to exclude (opt-outs, customers, etc.)')
    p.add_argument('--verifier', choices=['none', 'hunter'], default='none')
    p.add_argument('--approved-addresses', help='file listing the exact addresses approved for the verifier')
    p.add_argument('--max-verifications', type=int, default=20)
    p.add_argument('--max-rows', type=int, default=500)
    p.add_argument('--workers', type=int, default=6)
    a = p.parse_args(argv)

    raw = read_table(a.input)
    blockers = []
    if len(raw) > a.max_rows:
        blockers.append(f'input truncated to first {a.max_rows} rows (of {len(raw)}); raise --max-rows to process more')
        raw = raw[:a.max_rows]
    resolver = make_resolver()
    preflight = resolver.status('gmail.com')
    if preflight != 'mx_present':
        blockers.append('DNS lookups are not working in this environment, so no row can be SURE. '
                        'Enable network access (or run locally) and re-run.')
    inner = HunterVerifier() if a.verifier == 'hunter' else None
    approved = load_list(a.approved_addresses) if a.approved_addresses else []
    gate = VerifierGate(inner, approved, a.max_verifications)
    suppression = load_list(a.suppress) if a.suppress else []
    judge = Judge(SafeFetcher(), resolver, gate, crawl=not a.no_crawl, max_pages=a.max_pages,
                  evidence_policy=a.evidence_policy, suppression=suppression, workers=a.workers)
    started = now()
    results = judge.run(raw)
    meta = {'version': VERSION, 'input': str(a.input), 'started_at': started, 'finished_at': now(),
            'resolver': getattr(resolver, 'name', 'custom'), 'dns_preflight': preflight,
            'crawl': not a.no_crawl, 'max_pages': a.max_pages, 'evidence_policy': a.evidence_policy,
            'verifier': a.verifier, 'suppression': f'{len(suppression)} entries' if a.suppress else 'not provided',
            'blockers': blockers}
    groups = write_outputs(results, a.out, meta, gate)
    print(json.dumps({'rows': len(results), 'sure': len(groups['SURE']), 'review': len(groups['REVIEW']),
                      'reject': len(groups['REJECT']), 'out': str(Path(a.out).resolve()),
                      'blockers': blockers}, indent=2))


if __name__ == '__main__':
    main()
